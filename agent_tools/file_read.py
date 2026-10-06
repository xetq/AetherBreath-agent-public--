"""
工具名称: read_file
功能: 读取文件内容或列出文件夹内容，支持多种文档格式
"""

import logging
from pathlib import Path
from datetime import datetime
from typing import Optional

# 尝试导入项目配置（如果可用）
try:
    from agent import PROJECT_ROOT, WORKSPACE_ROOT
except ImportError:
    # 降级方案：使用当前工作目录作为工作空间
    WORKSPACE_ROOT = Path.cwd()
    PROJECT_ROOT = Path.cwd()

logger = logging.getLogger(__name__)

# ------------------------------------------------------------
# 辅助函数：检查路径是否在工作空间内（已注释限制）
# ------------------------------------------------------------
def _is_within_workspace(path: Path) -> bool:
    """检查路径是否在工作空间内（当前已无条件返回 True）"""
    return True  # 取消限制，允许访问任意路径

# ------------------------------------------------------------
# 文件类型判定（2026-09-23：主人要求支持 .tsx / .ts）
# ------------------------------------------------------------
# 明面上的文本类（含全部常见代码/配置后缀）。**代码类必须齐全** ——
# 只补 ts/tsx 的话，下次轮到 .jsx/.vue/.svelte 又得改一遍（同一族问题一次修完）。
TEXT_EXTS = {
    # 文档 / 数据 / 配置
    '.txt', '.md', '.markdown', '.rst', '.json', '.jsonl', '.csv', '.tsv', '.log',
    '.ini', '.cfg', '.conf', '.yaml', '.yml', '.toml', '.env', '.properties', '.lock',
    # 前端 / 脚本 / 样式
    '.js', '.jsx', '.ts', '.tsx', '.mjs', '.cjs', '.vue', '.svelte', '.astro',
    '.html', '.htm', '.css', '.scss', '.less', '.sass', '.styl',
    '.sh', '.bash', '.zsh', '.fish', '.bat', '.cmd', '.ps1', '.psm1',
    # 其它语言 / 查询
    '.py', '.pyi', '.pyw', '.go', '.rs', '.java', '.kt', '.kts', '.scala', '.swift',
    '.c', '.h', '.cc', '.cpp', '.cxx', '.hpp', '.hh', '.m', '.mm', '.cs',
    '.rb', '.php', '.lua', '.r', '.pl', '.pm', '.dart', '.ex', '.exs', '.erl',
    '.sql', '.graphql', '.gql', '.proto', '.gradle', '.cmake', '.mk', '.make',
    '.xml', '.svg', '.diff', '.patch',
}

# 有专用读取器的文档格式（不当文本读）
DOC_EXTS = {'.docx', '.pdf', '.xlsx'}

MAX_TEXT_BYTES = 10 * 1024 * 1024          # 纯文本上限 10MB（与旧版一致）


def _looks_like_text(path: Path, sniff: int = 8192) -> bool:
    """未知扩展名（含无扩展名）的兜底判定：取前 8KB，含 NUL 一律当二进制；
    再试 UTF-8 / GBK 能不能解（Dockerfile、Makefile、LICENSE 这类不该被挡在门外）。

    截断点可能落在多字节字符中间，所以整体解不开时再退 3 字节重试一次。
    """
    try:
        with open(path, 'rb') as f:
            head = f.read(sniff)
    except Exception:
        return False
    if not head:
        return True
    if b'\x00' in head:
        return False
    for enc in ('utf-8', 'gbk'):
        for cut in (head, head[:-3] if len(head) > 8 else head):
            try:
                cut.decode(enc)
                return True
            except UnicodeDecodeError:
                continue
    return False


def _read_text(path: Path):
    """按编码依次尝试读取：UTF-8（含 BOM）→ GBK（中文 Windows 常见）。
    返回 (文本, 实际编码)；全解不开返回 (None, '')。"""
    for enc in ('utf-8-sig', 'utf-8', 'gbk'):
        try:
            with open(path, 'r', encoding=enc) as f:
                return f.read(), enc
        except UnicodeDecodeError:
            continue
    return None, ''


# ------------------------------------------------------------
# 行窗口（offset / limit）
# ------------------------------------------------------------
def _apply_line_window(text: str, offset: Optional[int], limit: Optional[int]) -> str:
    """按行截取文件内容。offset 从 1 开始计数，limit 是要取的行数。

    为什么需要它：`web_extract` 抓超长页面时会截断成「头 75% + 尾 25%」，并在
    footer 里让模型用 `read_file path=... offset=N limit=200` 去补读被省略的中段。
    旧版 read_file 只接受 file_path，那条指引必然 `TypeError: unexpected keyword
    argument 'offset'` —— 截断之后唯一的补读路径是坏的（工具集审计 B3）。
    """
    if offset is None and limit is None:
        return text

    lines = text.splitlines()
    total = len(lines)

    try:
        start = int(offset) if offset is not None else 1
    except (TypeError, ValueError):
        start = 1
    if start < 1:
        start = 1
    if total and start > total:
        return f"⚠️ offset={start} 超出文件范围（该文件共 {total} 行）"

    end = total
    if limit is not None:
        try:
            want = int(limit)
        except (TypeError, ValueError):
            want = None
        if want is not None and want > 0:
            end = min(total, start + want - 1)

    body = "\n".join(lines[start - 1:end])
    if start == 1 and end >= total:
        return body
    head = f"[第 {start}-{end} 行 / 共 {total} 行"
    if end < total:
        head += f"；继续读：offset={end + 1} limit={max(1, end - start + 1)}"
    return head + "]\n" + body

# ------------------------------------------------------------
# 目录列表格式化
# ------------------------------------------------------------
def _format_directory_listing(directory: Path) -> str:
    """格式化目录列表，返回结构化文本"""
    items = []
    try:
        for item in sorted(directory.iterdir()):
            try:
                stat = item.stat()
                is_dir = item.is_dir()
                size = stat.st_size if not is_dir else 0
                mtime = datetime.fromtimestamp(stat.st_mtime).strftime('%Y-%m-%d %H:%M:%S')
                item_type = "📁" if is_dir else "📄"
                items.append(f"{item_type} {item.name}  {size} bytes  {mtime}")
            except PermissionError:
                items.append(f"🔒 {item.name} (无权限访问)")
        if not items:
            return "目录为空"
        return "\n".join(items)
    except PermissionError:
        return "❌ 无权限访问该目录"
    except Exception as e:
        return f"❌ 列出目录失败: {str(e)}"

# ------------------------------------------------------------
# 文档读取专用函数
# ------------------------------------------------------------
def _read_docx(file_path: Path) -> str:
    """读取 .docx 文件内容"""
    try:
        from docx import Document
    except ImportError:
        return "❌ 缺少依赖库 python-docx,请执行: pip install python-docx"
    try:
        doc = Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        return "\n".join(full_text)
    except Exception as e:
        return f"❌ 读取 DOCX 文件失败: {str(e)}"

def _read_pdf(file_path: Path) -> str:
    """读取 .pdf 文件内容（提取文本）"""
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            return "❌ 缺少依赖库 pypdf 或 PyPDF2,请执行: pip install pypdf"
    try:
        reader = PdfReader(file_path)
        full_text = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                full_text.append(text)
        return "\n".join(full_text) if full_text else "(PDF 中未提取到文本)"
    except Exception as e:
        return f"❌ 读取 PDF 文件失败: {str(e)}"

def _read_xlsx(file_path: Path) -> str:
    """读取 .xlsx 文件内容（转为文本表格）"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "❌ 缺少依赖库 openpyxl，请执行: pip install openpyxl"
    try:
        wb = load_workbook(file_path, data_only=True)
        all_text = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            all_text.append(f"=== 工作表: {sheet_name} ===")
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
                if row_text.strip():
                    all_text.append(row_text)
        return "\n".join(all_text) if all_text else "（XLSX 文件为空）"
    except Exception as e:
        return f"❌ 读取 XLSX 文件失败: {str(e)}"

# ------------------------------------------------------------
# 主函数
# ------------------------------------------------------------
def read_file(file_path: str, offset: Optional[int] = None,
              limit: Optional[int] = None) -> str:
    """
    读取文件内容或列出目录内容。

    参数:
        file_path: 文件或目录的路径（相对路径或绝对路径）
        offset: 可选，从第几行开始返回（从 1 开始计数）。仅对纯文本生效。
        limit: 可选，最多返回多少行。仅对纯文本生效。
               （offset/limit 由 web_extract 的截断 footer 指引使用：
                 长页面被截断后，用它补读被省略的中段）

    返回:
        - 如果是目录：列出目录下的文件和子目录
        - 如果是文件：根据扩展名选择读取方式
          * 纯文本与常见代码/配置后缀 → 文本读取（可用 offset/limit 取行窗口）
            （.ts .tsx .js .jsx .vue .css .html .sh .go .rs … 都算；无扩展名的文本文件也会探测后读取）
          * .docx → 提取段落文本
          * .pdf → 提取页面文字
          * .xlsx → 转为表格文本
    """
    # 1. 解析路径
    path = Path(file_path)
    if not path.is_absolute():
        path = WORKSPACE_ROOT / path

    # 2. 安全检查（已取消限制）
    # if not _is_within_workspace(path): ...

    # 3. 检查是否存在
    if not path.exists():
        return f"❌ 错误：路径 '{path}' 不存在。"

    # 4. 处理目录
    if path.is_dir():
        return _format_directory_listing(path)

    # 5. 处理文件（按扩展名分派；未知扩展名走"像不像文本"的兜底判定）
    ext = path.suffix.lower()
    try:
        # 有专用读取器的文档格式
        if ext in DOC_EXTS:
            return {'.docx': _read_docx, '.pdf': _read_pdf, '.xlsx': _read_xlsx}[ext](path)

        # 文本类：白名单直读；白名单外先探测（无扩展名的 Dockerfile / Makefile 也算）
        if ext in TEXT_EXTS or _looks_like_text(path):
            size = path.stat().st_size
            if size > MAX_TEXT_BYTES:
                return f"⚠️ 文件过大(超过 {MAX_TEXT_BYTES//1024//1024}MB),拒绝读取。"
            text, enc = _read_text(path)
            if text is None:
                return (f"❌ 读取失败：{ext or '（无扩展名）'} 文件既不是 UTF-8 也不是 GBK 文本，"
                        f"且不属于已知文档格式。")
            out = _apply_line_window(text, offset, limit)
            # 非 UTF-8 时标明编码：免得模型把 GBK 认成乱码去猜
            return out if enc.startswith('utf-8') else f"（编码 {enc}）\n{out}"

        # 二进制 / 不认识：说清楚支持什么，别只说一句"不支持"
        return (f"❌ 不支持读取该文件格式（{ext or '无扩展名'}）。\n"
                f"支持：纯文本与常见代码/配置后缀（.txt .md .py .js .ts .tsx .json .csv .yaml 等）、"
                f"无扩展名的文本文件（探测判定）、以及 .docx / .pdf / .xlsx。")
    except Exception as e:
        logger.error(f"读取文件失败: {e}")
        return f"❌ 读取文件出错: {str(e)}"

# ------------------------------------------------------------
# Schema（工具说明书）
# ------------------------------------------------------------
read_file_schema = {
    "type": "function",
    "function": {
        "name": "read_file",
        "description": (
            "读取文件内容或列出目录内容。\n"
            "支持的文件类型：\n"
            "  - 纯文本与代码/配置：.txt .md .json .csv .log .ini .yaml .env\n"
            "    .py .js .jsx .ts .tsx .vue .svelte .css .scss .html .xml .sh .bat .ps1\n"
            "    .sql .go .rs .java .kt .swift .c .h .cpp .cs .rb .php .lua .r .dart 等\n"
            "  - 无扩展名的文本文件（Dockerfile / Makefile / LICENSE）：按内容探测后读取\n"
            "  - 编码：UTF-8（含 BOM）优先，其次 GBK（会标明实际编码）\n"
            "  - Word：.docx\n"
            "  - PDF：.pdf\n"
            "  - Excel：.xlsx\n"
            "如果是目录，返回文件列表（含大小和修改时间）。\n"
            "纯文本可用 offset/limit 取行窗口（web_extract 截断后的补读就靠它）。"
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "file_path": {
                    "type": "string",
                    "description": "要读取的文件或目录路径（绝对或相对路径）"
                },
                "offset": {
                    "type": "integer",
                    "description": "可选，从第几行开始返回（从 1 开始计数）。仅对纯文本生效。"
                },
                "limit": {
                    "type": "integer",
                    "description": "可选，最多返回多少行。仅对纯文本生效。"
                }
            },
            "required": ["file_path"]
        }
    }
}
